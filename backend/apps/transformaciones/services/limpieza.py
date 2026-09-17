"""
Servicio de limpieza de datos (determinista, sin IA).

Lee el Excel origen y lo normaliza. Detecta el formato automáticamente:
  - TABLA: encabezados arriba, datos en filas.
  - VERTICAL: formulario "etiqueta | dato".

Y puede leer UNA hoja o TODAS las hojas del origen combinándolas (multi-hoja),
para orígenes con la información repartida en varias pestañas.

Todo aquí es reproducible: mismos datos de entrada → mismo resultado.
"""
import re
from decimal import Decimal, InvalidOperation

import pandas as pd
from openpyxl import load_workbook


def _a_numero(valor):
    """Convierte a número entendiendo formato chileno (1.240,50 → 1240.50)."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    s = str(valor).strip()
    if not s:
        return None
    s = re.sub(r"[^\d.,\-]", "", s)
    if not s:
        return None
    if "." in s and "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    elif "." in s:
        partes = s.split(".")
        if len(partes) > 2:
            s = s.replace(".", "")
        elif len(partes[1]) == 3:
            s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def normalizar_unidad(texto, catalogo):
    """Resuelve el texto de una unidad ('M3', 'mt2') a su símbolo canónico."""
    if not texto:
        return None
    t = str(texto).strip().lower()
    for unidad in catalogo:
        if t == unidad["simbolo"].lower():
            return unidad["simbolo"]
        if t in [a.lower() for a in unidad.get("alias", [])]:
            return unidad["simbolo"]
    return None


# ─────────────────────────────────────────────────────────────
# DETECCIÓN Y LECTURA DE FORMATO (tabla vs vertical, 1 hoja o todas)
# ─────────────────────────────────────────────────────────────

def _texto(c):
    return "" if c is None else str(c).strip()


def _detectar_formato(ws):
    """Decide si una hoja es TABLA o VERTICAL."""
    filas = [f for f in ws.iter_rows(values_only=True) if any(_texto(c) for c in f)]
    if not filas:
        return "vacia"
    ancho = max(sum(1 for c in f if _texto(c)) for f in filas)
    filas_par = sum(1 for f in filas if len([c for c in f if _texto(c)]) == 2)
    if ancho <= 3 and filas_par >= max(3, len(filas) * 0.6):
        return "vertical"
    return "tabla"


def _leer_vertical(ws):
    """Convierte una hoja vertical (etiqueta | dato) en un DataFrame de 1 fila."""
    datos = {}
    for fila in ws.iter_rows(values_only=True):
        celdas = [_texto(c) for c in fila if _texto(c)]
        if len(celdas) >= 2:
            etiqueta = celdas[0].rstrip(":").strip()
            valor = celdas[1]
            if etiqueta and etiqueta not in datos:
                datos[etiqueta] = valor
    return pd.DataFrame([datos]) if datos else pd.DataFrame()


def _leer_una_hoja(ruta_archivo, hoja):
    """Lee una hoja concreta detectando su formato. Devuelve (df, tipo)."""
    try:
        wb = load_workbook(ruta_archivo, data_only=True)
        ws = wb[hoja] if (hoja and hoja in wb.sheetnames) else wb.active
        tipo = _detectar_formato(ws)
        if tipo == "vertical":
            df = _leer_vertical(ws)
            wb.close()
            if not df.empty:
                return df, "vertical"
        wb.close()
    except Exception:
        pass
    df = pd.read_excel(ruta_archivo, sheet_name=hoja if hoja else 0, dtype=str)
    return df, "tabla"


def _leer_todas_las_hojas(ruta_archivo):
    """
    Lee TODAS las hojas y combina sus campos en un DataFrame de una fila.
    Si dos hojas tienen un campo con el mismo nombre, la segunda se prefija con
    el nombre de la hoja para no pisar a la primera.
    Solo tiene sentido cuando las hojas son verticales (formularios); si alguna
    es una tabla de varias filas, toma su primera fila.
    """
    wb = load_workbook(ruta_archivo, read_only=True)
    nombres = wb.sheetnames
    wb.close()

    combinado = {}
    for hoja in nombres:
        df, _tipo = _leer_una_hoja(ruta_archivo, hoja)
        if df.empty:
            continue
        fila = df.iloc[0]
        for col in df.columns:
            clave = str(col)
            if clave in combinado:
                clave = f"{hoja} · {col}"
            combinado[clave] = fila[col]

    return pd.DataFrame([combinado]) if combinado else pd.DataFrame()


def limpiar_excel(ruta_archivo, opciones, catalogo_unidades=None, hoja=None,
                  todas_las_hojas=False):
    """
    Lee y limpia el Excel.
      - Si todas_las_hojas=True: combina los campos de TODAS las hojas.
      - Si hoja se indica: lee esa hoja.
      - Si no: lee la hoja activa.
    Detecta tabla vs vertical automáticamente. Devuelve (dataframe, resumen).
    """
    catalogo_unidades = catalogo_unidades or []
    resumen = {
        "filas_originales": 0,
        "formato_detectado": "tabla",
        "hojas_leidas": 1,
        "duplicados_eliminados": 0,
        "espacios_corregidos": 0,
        "unidades_normalizadas": 0,
        "montos_convertidos": 0,
        "celdas_sin_interpretar": 0,
    }

    if todas_las_hojas:
        df = _leer_todas_las_hojas(ruta_archivo)
        resumen["formato_detectado"] = "multi-hoja"
        try:
            wb = load_workbook(ruta_archivo, read_only=True)
            resumen["hojas_leidas"] = len(wb.sheetnames)
            wb.close()
        except Exception:
            pass
    else:
        df, tipo = _leer_una_hoja(ruta_archivo, hoja=hoja)
        resumen["formato_detectado"] = tipo

    resumen["filas_originales"] = len(df)

    if df.empty:
        return df, resumen

    # 1. Espacios sobrantes.
    if opciones.get("espacios", True):
        df.columns = [str(c).strip() for c in df.columns]
        for col in df.select_dtypes(include="object").columns:
            antes = df[col].copy()
            df[col] = df[col].apply(lambda x: str(x).strip() if pd.notna(x) else x)
            resumen["espacios_corregidos"] += int((antes != df[col]).sum())

    # 2. Duplicados.
    if opciones.get("duplicados", True):
        n_antes = len(df)
        df = df.drop_duplicates().reset_index(drop=True)
        resumen["duplicados_eliminados"] = n_antes - len(df)

    # 3. Filas vacías.
    if opciones.get("vacias", False):
        df = df.dropna(how="all").reset_index(drop=True)

    return df, resumen


def detectar_columnas_numericas(df):
    """Indica qué columnas parecen numéricas."""
    columnas = []
    for col in df.columns:
        muestra = df[col].dropna().head(20)
        if len(muestra) == 0:
            continue
        convertibles = sum(1 for v in muestra if _a_numero(v) is not None)
        if convertibles / len(muestra) > 0.7:
            columnas.append(col)
    return columnas