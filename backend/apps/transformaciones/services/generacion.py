"""
Servicio de generación del documento final.

Maneja DOS tipos de hoja en la plantilla destino:
  1. FORMULARIOS verticales (etiqueta | casilla): escribe cada campo mapeado.
  2. ITEMIZADOS (tablas con encabezados): vuelca las partidas del origen,
     emparejando columnas por nombre y limpiando cada celda según su columna.

Preserva el formato del Excel original (colores, títulos, estilos).
"""
import io
import re
import unicodedata
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils import coordinate_to_tuple, get_column_letter

from .normalizacion import fecha, numero_en_palabras, resuelto_por_reglas, rut, telefono


FORMATOS_EXCEL = {
    "ENTERO":     ("#,##0", False),
    "DECIMAL":    ("#,##0.00", False),
    "PESOS":      ('"$"#,##0', False),
    "PESOS_DEC":  ('"$"#,##0.00', False),
    "PORCENTAJE": ("0.0%", True),
    "UF":         ('#,##0.00" UF"', False),
}


def _normalizar(texto):
    if texto is None:
        return ""
    s = str(texto).strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    s = "".join(c if c.isalnum() else " " for c in s)
    return " ".join(s.split())


def _a_numero(valor):
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    s = str(valor).strip()
    if not s:
        return None
    s = re.sub(r"\([^)]*\)", "", s)
    s = re.sub(r"\.?-\s*$", "", s.strip())
    s = re.sub(r"[^\d.,]", "", s)
    if not s:
        return None
    if "." in s and "," in s:
        if s.rfind(",") < s.rfind("."):
            s = s.replace(",", "")
        else:
            s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        if re.fullmatch(r"\d{1,3}(,\d{3})+", s):
            s = s.replace(",", "")
        else:
            s = s.replace(",", ".")
    elif "." in s:
        partes = s.split(".")
        # "0.125" es decimal: ningún número con miles empieza en 0.
        if partes[0] != "0" and all(len(g) == 3 for g in partes[1:]):
            s = s.replace(".", "")
        elif len(partes) > 2:
            s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


# ─────────────────────────────────────────────────────────────
# LIMPIEZA DETERMINISTA (Capa 1) — gratis, sin IA
# ─────────────────────────────────────────────────────────────

def _parece(nombre_o_etiqueta, palabras):
    txt = _normalizar(nombre_o_etiqueta)
    return any(p in txt for p in palabras)


def _limpiar_rut(v):
    formateado = rut(v)
    if formateado:
        return formateado
    s = str(v).strip().replace(" ", "")
    if re.match(r"^[\d.\-]+[kK\d]$", s):
        return s
    return v


def _limpiar_email(v):
    return str(v).replace(" ", "").strip()


def _limpiar_telefono(v):
    normal = telefono(v)
    if normal:
        return normal
    s = str(v)
    s = re.split(r"[/,]", s)[0].strip()
    s = re.sub(r"[^\d+]", "", s)
    return s if s else v


def _limpiar_monto(v):
    s = re.sub(r"\([^)]*\)", "", str(v))
    s = re.sub(r"[^\d.,]", "", s)
    if "," in s and "." in s:
        if s.rfind(",") < s.rfind("."):
            s = s.replace(",", "")
        else:
            s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        if re.fullmatch(r"\d{1,3}(,\d{3})+", s):
            s = s.replace(",", "")
        else:
            s = s.replace(",", ".")
    elif re.fullmatch(r"[1-9]\d{0,2}(\.\d{3})+\.?", s):
        # Solo puntos de miles ("1.500.000"). Un punto decimal ("5970672.7636",
        # como llegan los números de Excel) se conserva: quitarlo multiplicaba el monto.
        s = s.replace(".", "")
    s = s.rstrip(".-")
    try:
        num = float(s)
        return str(int(num)) if num == int(num) else str(num)
    except ValueError:
        return v


def _limpiar_numero(v):
    m = re.search(r"\d[\d.,]*", str(v))
    if not m:
        return v
    num = _a_numero(m.group().rstrip(".,"))
    if num is None:
        return m.group()
    return str(int(num)) if num == int(num) else str(num)


def _limpiar_cantidad(v):
    """Cantidad: extrae el número aunque venga con palabras: 'cien (100)' -> 100."""
    s = str(v)
    m = re.search(r"\(?\s*(\d[\d.,]*)\s*\)?", s)
    if m:
        num = m.group(1).replace(".", "").replace(",", ".")
        try:
            f = float(num)
            return str(int(f)) if f == int(f) else str(f)
        except ValueError:
            return m.group(1)
    return v


def _limpiar_nombre(v):
    s = re.sub(r"\([^)]*\)", "", str(v))
    s = re.sub(r"\bRUT\b.*", "", s, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", s).strip()


def _limpiar_texto(v):
    return re.sub(r"\s+", " ", str(v).strip())


_FORMATOS_MONTO = {"PESOS", "PESOS_DEC", "UF"}
_FORMATOS_NUMERO = {"ENTERO", "DECIMAL", "PORCENTAJE"}


def limpiar_valor(valor, campo):
    if valor is None:
        return valor
    s = str(valor).strip()
    if not s:
        return valor

    formato = getattr(campo, "formato_numero", "") or "NINGUNO"
    ref = f"{getattr(campo, 'nombre', '')} {getattr(campo, 'etiqueta_busqueda', '')}"

    if _parece(ref, ["rut", "rol unico"]):
        return _limpiar_rut(s)
    if _parece(ref, ["email", "correo", "mail"]):
        return _limpiar_email(s)
    if _parece(ref, ["telefono", "fono", "celular", "movil"]):
        return _limpiar_telefono(s)
    if _parece(ref, ["representante", "contacto", "persona"]):
        return _limpiar_nombre(s)

    if formato in _FORMATOS_MONTO:
        return _limpiar_monto(s)
    if formato in _FORMATOS_NUMERO:
        return _limpiar_numero(s)

    return _limpiar_texto(s)


# ─────────────────────────────────────────────────────────────
# ITEMIZADOS (tablas) — leer, emparejar, limpiar y volcar
# ─────────────────────────────────────────────────────────────

# Palabras que delatan el encabezado de una tabla de itemizado.
_PALABRAS_TABLA = [
    "descripcion", "cantidad", "precio", "valor", "total", "unidad",
    "cargo", "rol", "remuneracion", "item", "detalle", "personas", "monto", "pagar",
    "subtotal", "cant",
]


def _texto(c):
    return "" if c is None else str(c).strip()


# Límites de lectura. Hay Excel que "declaran" 1.048.576 filas (formato aplicado
# a columnas enteras): recorrerlas todas tardaba minutos. El encabezado se busca
# solo arriba, y la tabla termina tras varias filas seguidas vacías.
_MAX_FILAS_ENCABEZADO = 100
# Con 50 se leían como partidas los bloques que hay bajo la tabla (una lista de
# factores, notas…). Dos filas vacías seguidas ya separan otra cosa.
_MAX_FILAS_VACIAS = 2


TIPOS_HOJA = ("TABLA", "FORMULARIO")


def abrir_plantilla(ruta_plantilla, tipos_hoja=None, **kwargs):
    """
    Abre el Excel de una plantilla recordando qué hojas marcó el usuario a mano
    como "TABLA" o "FORMULARIO" (Plantilla.tipos_hoja). Así toda la detección de
    tablas sobre ese libro respeta la marca sin pasarla función por función.
    """
    wb = load_workbook(ruta_plantilla, **kwargs)
    wb._tipos_hoja = tipos_hoja or {}
    return wb


def _tipo_manual(ws):
    return getattr(getattr(ws, "parent", None), "_tipos_hoja", {}).get(ws.title, "")


def _es_dato(v):
    """Números y fechas no pueden ser encabezados de columna."""
    return isinstance(v, (int, float, date, datetime)) and not isinstance(v, bool)


def _tiene_palabra_tabla(texto):
    """
    ¿El texto tiene una palabra típica de encabezado? Por palabra completa (o
    su comienzo, para plurales: «Totales», «Precios»): buscada dentro del texto,
    «rol» aparecía en «Carol» y una fila de datos pasaba por encabezado.
    """
    return any(w.startswith(p) for w in _normalizar(texto).split() for p in _PALABRAS_TABLA)


def _encabezado_por_forma(filas):
    """
    Encabezado reconocido por la FORMA de la tabla, digan lo que digan sus
    textos: una fila con 3+ títulos distintos (sin números ni «:» de etiqueta)
    seguida de al menos 2 filas con datos en la mayoría de esas columnas.
    Solo sirve en el origen, donde la tabla trae filas llenas.
    """
    for i, fila in enumerate(filas):
        celdas = [_texto(c) for c in fila]
        cols = [j for j, c in enumerate(celdas) if c]
        if len(cols) < 3:
            continue
        textos = [celdas[j] for j in cols]
        if (any(_es_dato(fila[j]) for j in cols)
                or any(t.endswith(":") for t in textos)
                or len({_normalizar(t) for t in textos}) < len(textos)):
            continue
        minimo = (len(cols) + 1) // 2
        con_datos = sum(
            1 for f in filas[i + 1:i + 6]
            if sum(1 for j in cols if j < len(f) and _texto(f[j])) >= minimo
        )
        if con_datos >= 2:
            return i + 1, celdas
    return None


def _nombres_unicos(celdas):
    """
    Nombres de columna sin repetir: la segunda «SUB TOTAL» pasa a ser «SUB TOTAL (2)».
    Con dos tablas lado a lado (mismos títulos), los valores se guardan por
    nombre y la segunda tabla pisaba a la primera.
    """
    vistos, salida = {}, []
    for c in celdas:
        if c:
            clave = _normalizar(c)
            vistos[clave] = vistos.get(clave, 0) + 1
            if vistos[clave] > 1:
                c = f"{c} ({vistos[clave]})"
        salida.append(c)
    return salida


def _detectar_encabezado_tabla(ws, origen=False):
    r = _buscar_encabezado_tabla(ws, origen)
    return (r[0], _nombres_unicos(r[1])) if r else None


def _buscar_encabezado_tabla(ws, origen=False):
    """
    Busca la fila ENCABEZADO de una tabla. Devuelve (fila, [columnas]) o None.
      1. Marca manual de la plantilla: "FORMULARIO" = nunca es tabla.
      2. Palabras típicas: 3+ títulos y 2 de ellos como Cantidad, Precio, Total…
      3. Solo en el origen: la forma de la tabla (encabezado + filas con datos).
      4. Marca manual "TABLA": la primera fila con 2+ títulos es el encabezado.
    """
    tipo = _tipo_manual(ws)
    if tipo == "FORMULARIO":
        return None
    filas = list(ws.iter_rows(max_row=_MAX_FILAS_ENCABEZADO, values_only=True))
    for i, fila in enumerate(filas, start=1):
        celdas = [_texto(c) for c in fila]
        con_texto = [c for c in celdas if c]
        # Una fila con números o fechas es de datos, no de títulos.
        if len(con_texto) >= 3 and not any(_es_dato(c) for c in fila):
            coinc = sum(1 for c in con_texto if _tiene_palabra_tabla(c))
            if coinc >= 2:
                return i, celdas
    if origen:
        por_forma = _encabezado_por_forma(filas)
        if por_forma:
            return por_forma
    if tipo == "TABLA":
        for i, fila in enumerate(filas, start=1):
            celdas = [_texto(c) for c in fila]
            if sum(1 for c in celdas if c) >= 2:
                return i, celdas
    return None


def _leer_itemizado_origen(ws, ws_formulas=None):
    """
    Lee una tabla de partidas del origen. Devuelve dict o None.
    `ws` es la hoja con los valores calculados; `ws_formulas` (opcional), la
    misma hoja con sus fórmulas, para poder copiarlas al documento.
    """
    r = _detectar_encabezado_tabla(ws, origen=True)
    if not r:
        return None
    fila_enc, encabezados = r
    cols = [(idx, nom) for idx, nom in enumerate(encabezados) if nom]
    # `filas`: distancia de cada partida al encabezado, para escribirla en la
    # misma posición del destino (saltarse filas en 0 corría todo hacia arriba).
    partidas, filas = [], []
    vacias = 0
    for n, fila in enumerate(ws.iter_rows(min_row=fila_enc + 1, values_only=True), start=1):
        # Se guarda el valor ORIGINAL (fecha, número o texto), no su texto:
        # convertir todo a str perdía las fechas y los números.
        valores = {nom: (fila[idx] if idx < len(fila) else None) for idx, nom in cols}
        # Una fila cuenta si CUALQUIER columna trae dato: antes se exigía la
        # primera, y una primera columna vacía ("Columna1") descartaba la tabla.
        # Un 0 solo no es dato: las tablas de Excel arrastran fórmulas (Total =
        # Precio*Cantidad) hasta el final de la hoja, y en filas vacías dan 0.
        if any(_texto(v) and v != 0 for v in valores.values()):
            partidas.append(valores)
            filas.append(n)
            vacias = 0
        else:
            vacias += 1
            if vacias >= _MAX_FILAS_VACIAS:
                break
    if not partidas:
        return None
    datos = {"hoja": ws.title, "encabezados": [n for _, n in cols],
             "partidas": _normalizar_partidas(partidas), "filas": filas,
             "fila_encabezado": fila_enc, "indices": {nom: idx for idx, nom in cols}}
    if ws_formulas is not None:
        # Fórmulas de cada partida, {columna: "=E4*D4"}, en el mismo orden.
        con_formula = list(ws_formulas.iter_rows(
            min_row=fila_enc + 1, max_row=fila_enc + filas[-1], values_only=True))
        datos["formulas"] = [
            {nom: v for idx, nom in cols
             for v in [con_formula[n - 1][idx] if idx < len(con_formula[n - 1]) else None]
             if isinstance(v, str) and v.startswith("=")}
            for n in filas
        ]
    return datos


_REF_CELDA = re.compile(r"^\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?$|^\$?[A-Z]{1,3}:\$?[A-Z]{1,3}$|^\$?\d+:\$?\d+$")


def _formula_copiable(formula, ws_destino, se_mueve):
    """
    ¿Se puede copiar esta fórmula del origen a la hoja destino? Solo si todo lo
    que usa existe en la plantilla: las hojas que nombra y los nombres definidos
    (grati, tope…). Si la tabla cambia de lugar (`se_mueve`), las referencias a
    otras hojas no se pueden ajustar con seguridad: entonces tampoco.
    """
    from openpyxl.formula.tokenizer import Token, Tokenizer

    wb = ws_destino.parent
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return False
    for t in tokens:
        if t.type != Token.OPERAND or t.subtype != Token.RANGE:
            continue
        ref = t.value
        hoja = ws_destino
        if "!" in ref:
            nombre_hoja, ref = ref.rsplit("!", 1)
            nombre_hoja = nombre_hoja.strip("'").replace("''", "'")
            if se_mueve or nombre_hoja not in wb.sheetnames:
                return False
            hoja = wb[nombre_hoja]
        if "#REF" in ref or "[" in ref:
            return False      # referencia rota o a una tabla con nombre de Excel
        if _REF_CELDA.match(ref.upper()):
            continue
        # Es un nombre definido: tiene que existir en la plantilla.
        if ref not in hoja.defined_names and ref not in ws_destino.defined_names and ref not in wb.defined_names:
            return False
    return True


def _normalizar_partidas(partidas):
    """Limpia las partidas con pandas (fechas, tildes, espacios…), columna por columna."""
    import pandas as pd
    from .normalizacion import limpiar_dataframe

    df, _ = limpiar_dataframe(pd.DataFrame(partidas, dtype=object))
    df = df.astype(object).where(df.notna(), None)
    return df.to_dict("records")


def clave_ia_tabla(hoja, columna):
    """Clave con que se guardan los valores limpiados por IA de una columna de tabla."""
    return f"{hoja}::{columna}"


def columnas_ia_hoja(columnas_ia, hoja):
    """{columna destino: columna origen} que sugirió la IA para una hoja."""
    return {d: s["origen"] for d, s in ((columnas_ia or {}).get(hoja) or {}).items()}


def columnas_usuario_hoja(resultado_limpieza, hoja):
    """{columna destino: columna origen ("" = dejarla vacía)} elegidas por el usuario."""
    return dict(((resultado_limpieza or {}).get("columnas_usuario") or {}).get(hoja) or {})


def parejas_columnas(cols_origen, cols_destino, sugeridas=None, usuario=None):
    """
    De qué columna del origen sale cada columna destino, y por qué:
    {destino: (origen o None, fuente)}, con fuente "usuario", "nombre", "ia" o "ninguna".
    Manda lo que eligió el usuario (también "" = dejarla vacía); si no eligió,
    la del mismo nombre o parecido; si no hay, la que sugirió la IA.
    """
    por_nombre = _emparejar_por_nombre(cols_origen, cols_destino)
    usuario, sugeridas = usuario or {}, sugeridas or {}
    parejas = {}
    for d in cols_destino:
        if d in usuario:
            parejas[d] = (usuario[d] if usuario[d] in cols_origen else None, "usuario")
        elif por_nombre[d]:
            parejas[d] = (por_nombre[d], "nombre")
        elif sugeridas.get(d) in cols_origen:
            parejas[d] = (sugeridas[d], "ia")
        else:
            parejas[d] = (None, "ninguna")
    return parejas


def _emparejar_columnas(cols_origen, cols_destino, sugeridas=None, usuario=None):
    """{columna destino: columna origen o None}. Ver parejas_columnas."""
    return {d: o for d, (o, _) in parejas_columnas(cols_origen, cols_destino, sugeridas, usuario).items()}


def _emparejar_por_nombre(cols_origen, cols_destino):
    mapa = {}
    for col_dest in cols_destino:
        nd = set(_normalizar(col_dest).split())
        mejor, mejor_score = None, None
        for col_ori in cols_origen:
            no = set(_normalizar(col_ori).split())
            comunes = nd & no
            if not comunes:
                continue
            score = (_normalizar(col_ori) == _normalizar(col_dest), len(comunes), -len(no - nd))
            if mejor_score is None or score > mejor_score:
                mejor, mejor_score = col_ori, score
        mapa[col_dest] = mejor
    return mapa


# Tipos de columna, reconocidos por palabras de su encabezado.
_COL_DINERO = ["valor", "precio", "total", "remuneracion", "monto", "pagar", "venta",
               "costo", "neto", "iva", "subtotal", "importe", "sueldo", "pago"]
_COL_CANTIDAD = ["cantidad", "cant", "personas", "unidades"]
_COL_FECHA = ["fecha", "inicio", "termino", "vencimiento", "emision"]

# Errores de fórmula de Excel: no son datos, la celda queda vacía.
_ERRORES_EXCEL = {"#VALUE!", "#DIV/0!", "#N/A", "#REF!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}

_FORMATOS_FECHA_TEXTO = ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y",
                         "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M")


def _tipo_columna(nombre_columna):
    n = _normalizar(nombre_columna)
    palabras = n.split()
    if any(p in n for p in _COL_FECHA):
        return "fecha"
    if palabras and palabras[0] in _COL_CANTIDAD:   # "cantidad total" es cantidad
        return "cantidad"
    if any(p in n for p in _COL_DINERO):
        return "dinero"
    if any(p in palabras or n.startswith(p) for p in _COL_CANTIDAD):
        return "cantidad"
    return "texto"


def _entero_si_exacto(num):
    return int(num) if float(num).is_integer() else num


def _fecha_desde_texto(s):
    for fmt in _FORMATOS_FECHA_TEXTO:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _limpiar_celda_tabla(valor, nombre_columna):
    """
    Limpia una celda de itemizado y devuelve el valor con su tipo real
    (date, int/float o str), o None si queda vacía.
      - Errores de Excel (#VALUE!, #DIV/0!…) → vacía.
      - Fechas → date (sin la hora 00:00:00).
      - Columnas de dinero/cantidad → número, aunque venga con símbolos ("1.500+", "$ 2.000").
      - Texto que es un número limpio ("100") → número.
      - Resto → texto sin espacios de más.
    """
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date() if (valor.hour, valor.minute, valor.second) == (0, 0, 0) else valor
    if isinstance(valor, date):
        return valor
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, (int, float)):
        return _entero_si_exacto(valor)

    s = re.sub(r"\s+", " ", str(valor)).strip()
    if not s or s.upper() in _ERRORES_EXCEL:
        return None

    tipo = _tipo_columna(nombre_columna)
    if tipo == "fecha" or re.fullmatch(r"\d{2}-\d{2}-\d{4}", s):
        f = _fecha_desde_texto(s)
        if f:
            return f
    if tipo in ("dinero", "cantidad"):
        # En cantidades se toma el número aunque venga con palabras: "cien (100)"
        # → 100, y también "cien" a secas.
        num = _a_numero(_limpiar_cantidad(s) if tipo == "cantidad" else s)
        if num is None and tipo == "cantidad":
            num = numero_en_palabras(s)
        if num is not None:
            return _entero_si_exacto(num)
        return s
    # Número escrito como texto ("100", "2,5"); sin ceros a la izquierda para no
    # romper códigos como "007".
    if re.fullmatch(r"-?(0|[1-9]\d*)([.,]\d+)?", s):
        return _entero_si_exacto(float(s.replace(",", ".")))
    return s


def texto_celda_tabla(valor, nombre_columna):
    """Cómo se ve una celda limpia en la vista previa (formato chileno)."""
    v = _limpiar_celda_tabla(valor, nombre_columna)
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d-%m-%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        miles = f"{v:,.0f}" if float(v).is_integer() else f"{v:,.2f}"
        miles = miles.replace(",", "X").replace(".", ",").replace("X", ".")
        return f"${miles}" if _tipo_columna(nombre_columna) == "dinero" else miles
    return str(v)


def _miles_cl(num, decimales):
    """1234567.891 → «1.234.567,89» (formato chileno)."""
    texto = f"{num:,.{decimales}f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def texto_campo(valor, campo):
    """
    Cómo se verá un campo de formulario en el documento, para la vista previa.
    Mostrar el número crudo («5970672.7636») confundía: con el punto chileno de
    miles se lee como 59 mil millones.
    """
    limpio = limpiar_valor(valor, campo)
    if limpio is None:
        return None
    formato = getattr(campo, "formato_numero", "") or "NINGUNO"
    es_numero = formato in FORMATOS_EXCEL or getattr(campo, "tipo", "") == "NUMERO"
    num = _a_numero(limpio) if es_numero else None
    if num is None:
        return str(limpio)
    if formato == "PESOS":
        return f"${_miles_cl(num, 0)}"
    if formato == "PESOS_DEC":
        return f"${_miles_cl(num, 2)}"
    if formato == "UF":
        return f"{_miles_cl(num, 2)} UF"
    if formato == "PORCENTAJE":
        return f"{_miles_cl(num, 1)}%"
    if formato == "ENTERO":
        return _miles_cl(num, 0)
    return _miles_cl(num, 0 if float(num).is_integer() else 2)


def _es_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def columnas_calculo_total(cols_dest):
    """
    Columnas para calcular total = precio × cantidad, si la tabla las tiene:
    (total, cantidad, precio) o None.
    """
    total = next((c for c in cols_dest
                  if "total" in _normalizar(c).split() and _tipo_columna(c) == "dinero"), None)
    cantidad = next((c for c in cols_dest if _tipo_columna(c) == "cantidad"), None)
    precio = next((c for c in cols_dest
                   if c != total and _tipo_columna(c) == "dinero"
                   and any(p in _normalizar(c) for p in ["precio", "unitario", "valor", "venta"])), None)
    return (total, cantidad, precio) if total and cantidad and precio else None


def filas_itemizado(datos_origen, cols_dest, calcular_totales=False, valores_ia=None,
                    columnas_ia=None, columnas_usuario=None):
    """
    Filas limpias de una tabla, listas para la vista previa o el documento.
    Devuelve (filas, faltantes): cada fila es (valores {columna: valor},
    columnas_calculadas, columnas_limpiadas_por_ia). `faltantes` = filas sin
    total pero con precio y cantidad; si `calcular_totales`, a esas se les pone
    precio × cantidad. `valores_ia` = {clave_ia_tabla: {original: limpio}}.
    `columnas_ia` = {columna destino: columna origen} sugeridas por la IA;
    `columnas_usuario` = las que eligió el usuario (mandan sobre todo).
    """
    mapa = _emparejar_columnas(datos_origen["encabezados"], cols_dest, columnas_ia, columnas_usuario)
    calc = columnas_calculo_total(cols_dest)
    valores_ia = valores_ia or {}
    hoja = datos_origen.get("hoja", "")
    filas, faltantes = [], 0
    for partida in datos_origen["partidas"]:
        valores, por_ia = {}, set()
        for col in cols_dest:
            ori = mapa.get(col)
            if not ori:
                valores[col] = None
                continue
            crudo = partida.get(ori)
            limpios = valores_ia.get(clave_ia_tabla(hoja, ori), {})
            if isinstance(crudo, str) and crudo.strip() in limpios:
                nuevo = limpios[crudo.strip()]
                if nuevo != crudo.strip():
                    por_ia.add(col)
                crudo = nuevo
            valores[col] = _limpiar_celda_tabla(crudo, col)
        calculadas = set()
        if calc:
            tot, cant, pre = calc
            if valores.get(tot) is None and _es_num(valores.get(cant)) and _es_num(valores.get(pre)):
                faltantes += 1
                if calcular_totales:
                    valores[tot] = _entero_si_exacto(round(valores[cant] * valores[pre], 2))
                    calculadas.add(tot)
        filas.append((valores, calculadas, por_ia))
    return filas, faltantes


def textos_para_ia(datos_origen, maximo=300):
    """
    Valores de texto de una tabla que vale la pena mandar a la IA: solo columnas
    de texto y fecha (nunca dinero ni cantidad), sin repetidos y solo los que la
    limpieza normal deja como texto. Devuelve [(clave, columna, tipo, valor)].
    """
    hoja = datos_origen.get("hoja", "")
    vistos, salida = set(), []
    for col in datos_origen["encabezados"]:
        tipo = _tipo_columna(col)
        if tipo in ("dinero", "cantidad"):
            continue   # privacidad: montos y cantidades no salen de la empresa
        for partida in datos_origen["partidas"]:
            crudo = partida.get(col)
            # Menos de 3 caracteres ("r", "da"): restos de celdas, no texto que corregir.
            if not isinstance(crudo, str) or len(crudo.strip()) < 3:
                continue
            if not isinstance(_limpiar_celda_tabla(crudo, col), str) or resuelto_por_reglas(crudo):
                continue   # ya quedó como número, fecha, RUT…: no hace falta la IA
            clave = clave_ia_tabla(hoja, col)
            if (clave, crudo.strip()) in vistos:
                continue
            vistos.add((clave, crudo.strip()))
            salida.append((clave, col, "FECHA" if tipo == "fecha" else "TEXTO", crudo.strip()))
            if len(salida) >= maximo:
                return salida
    return salida


def _volcar_itemizado(ws_destino, datos_origen, calcular_totales=False, valores_ia=None,
                      columnas_ia=None, columnas_usuario=None):
    """
    Escribe las partidas del origen en la tabla del destino, emparejando columnas
    y limpiando cada celda. Devuelve cuántas partidas escribió.
    """
    r = _detectar_encabezado_tabla(ws_destino)
    if not r:
        return 0
    fila_enc, encabezados_dest = r
    cols_dest = [(idx, nom) for idx, nom in enumerate(encabezados_dest) if nom]
    nombres_dest = [n for _, n in cols_dest]

    filas, _ = filas_itemizado(datos_origen, nombres_dest, calcular_totales, valores_ia,
                               columnas_ia, columnas_usuario)
    # Cada partida va a la misma distancia del encabezado que en el origen.
    posiciones = datos_origen.get("filas") or range(1, len(filas) + 1)

    # Fórmulas del origen: se copian (ajustadas, como al copiar y pegar en Excel)
    # solo si la tabla tiene la misma forma, es decir, si todas las columnas se
    # corren igual. Si no, una fórmula apuntaría a otra columna: va el valor.
    mapa = _emparejar_columnas(datos_origen["encabezados"], nombres_dest, columnas_ia, columnas_usuario)
    formulas = datos_origen.get("formulas")
    indices = datos_origen.get("indices") or {}
    corrimientos = {idx - indices[mapa[nom]] for idx, nom in cols_dest if mapa.get(nom) in indices}
    if not formulas or len(corrimientos) != 1:
        formulas = None
    else:
        dx = corrimientos.pop()
        dy = fila_enc - datos_origen["fila_encabezado"]

    escritas = 0
    for k, ((valores, _calculadas, _por_ia), n) in enumerate(zip(filas, posiciones)):
        for idx_col, nom_dest in cols_dest:
            valor = valores.get(nom_dest)
            formula = formulas[k].get(mapa.get(nom_dest)) if formulas else None
            if formula:
                desde = f"{get_column_letter(indices[mapa[nom_dest]] + 1)}{datos_origen['fila_encabezado'] + n}"
                hacia = f"{get_column_letter(idx_col + 1)}{fila_enc + n}"
                if _formula_copiable(formula, ws_destino, (dx, dy) != (0, 0)):
                    try:
                        valor = Translator(formula, origin=desde).translate_formula(hacia)
                    except Exception:
                        pass   # no se pudo ajustar: queda el valor calculado
            if valor is None:
                continue
            celda = ws_destino.cell(row=fila_enc + n, column=idx_col + 1)
            # Una celda combinada no se puede escribir (solo su esquina), y una
            # fórmula de la plantilla se respeta: Excel la recalcula con los datos.
            if isinstance(celda, MergedCell) or (isinstance(celda.value, str) and celda.value.startswith("=")):
                continue
            celda.value = valor
            if isinstance(valor, str) and valor.startswith("="):
                continue   # fórmula: el formato es el de la plantilla
            if celda.number_format != "General":
                continue   # el formato que ya trae la plantilla manda
            es_numero = isinstance(valor, (int, float)) and not isinstance(valor, bool)
            if isinstance(valor, datetime):
                celda.number_format = "DD-MM-YYYY HH:MM"
            elif isinstance(valor, date):
                celda.number_format = "DD-MM-YYYY"
            elif es_numero and _tipo_columna(nom_dest) == "dinero":
                celda.number_format = '"$"#,##0' if float(valor).is_integer() else '"$"#,##0.00'
        escritas += 1
    return escritas


# ─────────────────────────────────────────────────────────────
# BÚSQUEDA DE CASILLA Y ESCRITURA (formularios verticales)
# ─────────────────────────────────────────────────────────────

def _buscar_etiqueta(ws, texto):
    RELLENO = {"de", "del", "la", "el", "los", "las", "y", "o", "a"}

    def palabras_clave(s):
        return {p for p in _normalizar(s).split() if p not in RELLENO}

    objetivo_norm = _normalizar(texto)
    objetivo_clave = palabras_clave(texto)
    if not objetivo_norm:
        return None

    mejor_puntaje = 0
    mejor_pos = None

    for fila in ws.iter_rows():
        for celda in fila:
            if celda.value is None:
                continue
            valor_norm = _normalizar(celda.value)
            if not valor_norm:
                continue
            if valor_norm == objetivo_norm:
                return celda.row, celda.column + 1
            valor_clave = palabras_clave(celda.value)
            comunes = objetivo_clave & valor_clave
            # Al menos la mitad de las palabras de la etiqueta: con una sola en
            # común, "Monto garantía extra" caía en la columna "Monto" de una tabla.
            if comunes and len(comunes) * 2 >= len(objetivo_clave):
                puntaje = len(comunes)
                if objetivo_clave and objetivo_clave <= valor_clave:
                    puntaje += 10
                if puntaje > mejor_puntaje:
                    mejor_puntaje = puntaje
                    mejor_pos = (celda.row, celda.column + 1)

    return mejor_pos


def hoja_del_campo(wb, campo):
    """Hoja donde va el campo: la indicada, o la activa si no tiene. None si la
    hoja indicada ya no existe en el Excel de la plantilla."""
    if not campo.hoja_destino:
        return wb.active
    return wb[campo.hoja_destino] if campo.hoja_destino in wb.sheetnames else None


def ubicar_campo(ws, campo, tabla=False):
    """
    Celda donde se escribe un campo de formulario. Devuelve ((fila, col), "") o
    (None, motivo) con el motivo en palabras del usuario.

    `tabla` es el encabezado de tabla de la hoja si ya se calculó (False = calcularlo).
    """
    if ws is None:
        return None, (f"su hoja «{campo.hoja_destino}» no existe en el Excel de la plantilla. "
                      "Asígnalo a otra hoja o elimínalo.")
    destino = None
    if campo.celda_destino:
        try:
            destino = coordinate_to_tuple(campo.celda_destino)
        except ValueError:
            return None, (f"la celda «{campo.celda_destino}» no es válida. Usa una "
                          "referencia como B4.")
    if destino is None and campo.etiqueta_busqueda:
        destino = _buscar_etiqueta(ws, campo.etiqueta_busqueda)
    if destino is None:
        destino = _buscar_etiqueta(ws, campo.nombre.replace("_", " "))
    if destino is None:
        buscado = campo.etiqueta_busqueda or campo.nombre.replace("_", " ")
        return None, (f"no se encontró el texto «{buscado}» en la hoja «{ws.title}». "
                      "Escribe en «Etiqueta de búsqueda» el texto exacto de la casilla, "
                      "o elimina el campo si no corresponde.")

    if tabla is False:
        tabla = _detectar_encabezado_tabla(ws)
    if tabla and destino[0] >= tabla[0]:
        return None, (f"cae dentro de la tabla de la hoja «{ws.title}». Las columnas de "
                      "una tabla se llenan automáticamente con la tabla del archivo de origen: "
                      "este campo sobra y se puede eliminar.")
    return destino, ""


def revisar_plantilla(ruta_plantilla, campos, tipos_hoja=None):
    """
    Revisión de la plantilla para el editor de campos. Por hoja: si es tabla
    (con sus columnas) o formulario; qué campos no se podrán escribir y por qué;
    y dónde se escribirá cada uno de los demás (celda y etiqueta encontrada, para
    que el usuario vea que no hace falta escribir el texto exacto de la casilla).
      {"hojas": [{nombre, tipo, columnas}], "problemas": {campo_id: motivo},
       "ubicaciones": {campo_id: {hoja, celda, etiqueta, exacta}}}
    """
    from openpyxl.utils import get_column_letter

    wb = abrir_plantilla(ruta_plantilla, tipos_hoja)
    tablas = {ws.title: _detectar_encabezado_tabla(ws) for ws in wb.worksheets}
    hojas = [
        {"nombre": nombre, "tipo": "tabla" if enc else "formulario",
         "columnas": [c for c in enc[1] if c] if enc else [],
         # "" = detectado automáticamente; "TABLA"/"FORMULARIO" = marcado a mano.
         "tipo_manual": (tipos_hoja or {}).get(nombre, "")}
        for nombre, enc in tablas.items()
    ]
    problemas, ubicaciones = {}, {}
    for campo in campos:
        ws = hoja_del_campo(wb, campo)
        destino, motivo = ubicar_campo(ws, campo, tablas.get(ws.title) if ws is not None else None)
        if motivo:
            problemas[str(campo.id)] = motivo
            continue
        fila, col = destino
        # La etiqueta está a la izquierda de la casilla (salvo celda fija).
        etiqueta = "" if campo.celda_destino or col < 2 else _texto(ws.cell(row=fila, column=col - 1).value)
        buscado = campo.etiqueta_busqueda or campo.nombre.replace("_", " ")
        ubicaciones[str(campo.id)] = {
            "hoja": ws.title,
            "celda": f"{get_column_letter(col)}{fila}",
            "etiqueta": etiqueta,
            "exacta": _normalizar(etiqueta) == _normalizar(buscado),
        }
    wb.close()
    return {"hojas": hojas, "problemas": problemas, "ubicaciones": ubicaciones}


def _escribir_valor(celda, valor, campo):
    valor = limpiar_valor(valor, campo)
    formato = getattr(campo, "formato_numero", "") or "NINGUNO"
    if getattr(campo, "tipo", "") == "FECHA":
        f = fecha(valor)
        if isinstance(f, date):
            celda.value = f
            celda.number_format = "DD-MM-YYYY"
            return
    if formato in FORMATOS_EXCEL:
        code, es_pct = FORMATOS_EXCEL[formato]
        num = _a_numero(valor)
        if num is not None:
            if es_pct:
                num = num / 100.0
            celda.value = num
            celda.number_format = code
            return
    if getattr(campo, "tipo", "") == "NUMERO":
        # Sin formato elegido, un número igual se escribe como número: como
        # texto («695») Excel no lo suma y las fórmulas que lo usan fallan.
        num = _a_numero(valor)
        if num is not None:
            celda.value = _entero_si_exacto(num)
            return
    celda.value = valor


def buscar_tabla_origen(wb_origen, nombre_hoja, wb_formulas=None):
    """Tabla del origen para una hoja de tabla de la plantilla: primero la hoja
    del mismo nombre; si no, la primera hoja del origen que tenga tabla.
    `wb_formulas`: el origen abierto con fórmulas, para poder copiarlas."""
    def leer(hoja):
        return _leer_itemizado_origen(wb_origen[hoja], wb_formulas[hoja] if wb_formulas else None)

    if nombre_hoja in wb_origen.sheetnames:
        datos = leer(nombre_hoja)
        if datos:
            return datos
    for hoja_ori in wb_origen.sheetnames:
        datos = leer(hoja_ori)
        if datos:
            return datos
    return None


AVISO_TABLAS_FALLARON = (
    "No se pudieron copiar las tablas del archivo de origen; las hojas de tabla "
    "pueden quedar vacías. Revisa que el origen sea un Excel válido y vuelve a intentarlo."
)


def aviso_tabla_fallo(nombre_hoja):
    return (f"No se pudo copiar la tabla de la hoja «{nombre_hoja}», así que puede haber "
            "quedado vacía o incompleta. Revisa esa hoja en el documento generado.")


def aviso_tabla_sin_origen(nombre_hoja):
    return (f"La hoja «{nombre_hoja}» es una tabla, pero el archivo de origen no tiene "
            "ninguna tabla reconocible para llenarla, así que quedará vacía. Revisa que "
            "el origen tenga una fila de encabezados (Descripción, Cantidad, Precio…).")


def estado_campos(ruta_plantilla, campos, valores, tipos_hoja=None):
    """
    Estado de cada campo de formulario en el documento: su hoja y, si quedará
    vacío, el aviso con el motivo ("" si se llenará). `valores` = {campo_id: valor}
    de los campos que tienen un dato asignado; un campo que no está ahí no
    recibió ningún dato del origen.
    """
    wb = abrir_plantilla(ruta_plantilla, tipos_hoja)
    tablas = {ws.title: _detectar_encabezado_tabla(ws) for ws in wb.worksheets}
    estados = []
    for campo in campos:
        valor = valores.get(campo.id)
        ws = hoja_del_campo(wb, campo)
        _, motivo = ubicar_campo(ws, campo, tablas.get(ws.title) if ws is not None else None)
        if motivo:
            aviso = f"«{campo.nombre}» no se escribirá: {motivo}"
        elif campo.id not in valores:
            aviso = (f"«{campo.nombre}» quedará vacío: no tiene ningún dato del origen "
                     "asignado. Asígnale uno en el mapeo de abajo.")
        elif valor is None or str(valor).strip() in ("", "nan"):
            aviso = (f"«{campo.nombre}» quedará vacío: el dato que tiene asignado "
                     "no trae valor en el archivo de origen.")
        else:
            aviso = ""
        hoja = ws.title if ws is not None else campo.hoja_destino
        estados.append({"campo": campo, "hoja": hoja, "aviso": aviso})
    wb.close()
    return estados


def campos_sin_escribir(ruta_plantilla, campos, valores, tipos_hoja=None):
    """
    Avisos de los campos de formulario que quedarán vacíos en el documento, con
    el motivo. Lo usan la vista previa y la generación, para que digan lo mismo.
    """
    return [e["aviso"] for e in estado_campos(ruta_plantilla, campos, valores, tipos_hoja) if e["aviso"]]


def pedido_ia_columnas(ruta_plantilla, ruta_origen, sin_pareja, tipos_hoja=None,
                       anonimizar=True, muestras=3):
    """
    Lo que se manda a la IA para emparejar columnas de tabla que no calzaron por
    nombre: por hoja, esas columnas destino y las columnas del origen con unos
    pocos ejemplos. Montos y cantidades nunca salen: van como «<número oculto>»
    (y con `anonimizar`, cualquier número).
      {hoja: {"titulo", "columnas_destino", "sin_pareja": [...], "origen": {columna: [ejemplos]}}}
    El título y las demás columnas de la plantilla dan el contexto: sin ellos,
    «Numero» junto a «Nombre» y «Correo» no se entendía como el teléfono.
    """
    pedido = {}
    wb_dest = abrir_plantilla(ruta_plantilla, tipos_hoja)
    wb_ori = load_workbook(ruta_origen, data_only=True, read_only=True)
    try:
        for hoja, columnas in sin_pareja.items():
            datos = buscar_tabla_origen(wb_ori, hoja)
            enc = _detectar_encabezado_tabla(wb_dest[hoja]) if hoja in wb_dest.sheetnames else None
            if not datos or not columnas or not enc:
                continue
            titulo = " ".join(
                _texto(c) for fila in wb_dest[hoja].iter_rows(max_row=enc[0] - 1, values_only=True)
                for c in fila if _texto(c))[:200]
            origen = {}
            for col in datos["encabezados"]:
                oculta = _tipo_columna(col) in ("dinero", "cantidad")
                ejemplos = []
                for partida in datos["partidas"]:
                    v = partida.get(col)
                    if v is None or not _texto(v):
                        continue
                    if oculta or (anonimizar and _es_num(v)):
                        ejemplos.append("<número oculto>")
                    else:
                        ejemplos.append(_texto(v)[:40])
                    if len(ejemplos) >= muestras:
                        break
                origen[col] = ejemplos
            pedido[hoja] = {"titulo": titulo, "columnas_destino": [c for c in enc[1] if c],
                            "sin_pareja": list(columnas), "origen": origen}
    finally:
        wb_dest.close()
        wb_ori.close()
    return pedido


def emparejamiento_tablas(ruta_plantilla, ruta_origen, tipos_hoja=None, columnas_ia=None,
                          columnas_usuario=None):
    """
    Qué tan bien se emparejan las columnas de cada tabla de la plantilla con las
    del origen. Por hoja: {columnas, emparejadas, sin_pareja, vacias_elegidas,
    por_ia, puntajes}. Cada puntaje es 100 (mismo nombre o elegida por el
    usuario), 70 (nombre parecido), el de la IA (máximo 60, para que se revise)
    o 0 (sin pareja). `por_ia` lista las parejas que sugirió la IA:
    [{destino, origen, motivo}]; `vacias_elegidas`, las que el usuario decidió
    dejar vacías (no son un aviso).
    """
    resultado = {}
    try:
        wb_dest = abrir_plantilla(ruta_plantilla, tipos_hoja)
        wb_ori = load_workbook(ruta_origen, data_only=True, read_only=True)
    except Exception:
        return resultado
    for nombre in wb_dest.sheetnames:
        enc = _detectar_encabezado_tabla(wb_dest[nombre])
        if not enc:
            continue
        datos = buscar_tabla_origen(wb_ori, nombre)
        if not datos:
            continue
        cols = [c for c in enc[1] if c]
        sugeridas = (columnas_ia or {}).get(nombre) or {}
        parejas = parejas_columnas(datos["encabezados"], cols, columnas_ia_hoja(columnas_ia, nombre),
                                   (columnas_usuario or {}).get(nombre))

        def puntaje(c):
            origen, fuente = parejas[c]
            if fuente == "usuario":
                return 100
            if not origen:
                return 0
            if fuente == "ia":
                return min(int(sugeridas[c].get("confianza") or 0), 60)
            return 100 if _normalizar(origen) == _normalizar(c) else 70

        resultado[nombre] = {
            "columnas": len(cols),
            "emparejadas": sum(1 for c in cols if parejas[c][0]),
            "sin_pareja": [c for c in cols if parejas[c] == (None, "ninguna")],
            "vacias_elegidas": [c for c in cols if parejas[c] == (None, "usuario")],
            "por_ia": [{"destino": c, "origen": parejas[c][0], "motivo": sugeridas[c].get("motivo", "")}
                       for c in cols if parejas[c][1] == "ia"],
            "puntajes": [puntaje(c) for c in cols],
        }
    wb_dest.close()
    wb_ori.close()
    return resultado


def generar_documento(ruta_plantilla, mapeos_con_valor, ruta_origen=None, calcular_totales=False,
                      valores_ia_tablas=None, avisos=None, tipos_hoja=None, columnas_ia=None,
                      columnas_usuario=None):
    """
    ruta_plantilla: .xlsx de la plantilla destino.
    mapeos_con_valor: lista de {campo, valor} para los campos de FORMULARIO.
    ruta_origen: (opcional) .xlsx del origen, para volcar los ITEMIZADOS.
                 Si se pasa, busca tablas en el origen y las vuelca a las hojas
                 del destino que tengan tabla.
    avisos: (opcional) lista donde se agrega, en palabras del usuario, todo lo
            que no se pudo escribir (campos sin valor o sin casilla, tablas vacías).
    """
    wb = abrir_plantilla(ruta_plantilla, tipos_hoja)
    if avisos is None:
        avisos = []

    # 1. Escribir los campos de formulario. Lo que no se puede escribir no se
    #    salta en silencio: queda en `avisos` para mostrárselo al usuario.
    for item in mapeos_con_valor:
        campo = item["campo"]
        valor = item["valor"]
        if valor is None:
            continue  # Sin dato: ya lo avisa campos_sin_escribir().
        ws = hoja_del_campo(wb, campo)
        destino, _ = ubicar_campo(ws, campo)
        if destino is None:
            continue  # Sin casilla: ya lo avisa campos_sin_escribir().
        fila, col = destino
        celda = ws.cell(row=fila, column=col)
        if isinstance(celda, MergedCell):
            # Casilla combinada: se escribe en su esquina superior izquierda.
            rango = next(r for r in ws.merged_cells.ranges if celda.coordinate in r)
            celda = ws.cell(row=rango.min_row, column=rango.min_col)
        _escribir_valor(celda, valor, campo)

    # 2. Volcar los ITEMIZADOS: para cada hoja del destino que tenga tabla,
    #    buscar en el origen una tabla (preferir la hoja del mismo nombre).
    if ruta_origen:
        try:
            wb_origen = load_workbook(ruta_origen, data_only=True, read_only=True)
        except Exception:
            wb_origen = None
            avisos.append(AVISO_TABLAS_FALLARON)
        try:
            # El mismo origen con sus fórmulas, para copiarlas a las tablas.
            wb_formulas = load_workbook(ruta_origen, read_only=True) if wb_origen else None
        except Exception:
            wb_formulas = None
        for nombre_hoja in (wb.sheetnames if wb_origen else []):
            ws_dest = wb[nombre_hoja]
            # Cada hoja por separado: antes un error en una tabla cortaba el
            # ciclo y todas las hojas siguientes quedaban vacías.
            try:
                # ¿La hoja destino tiene una tabla de itemizado?
                if _detectar_encabezado_tabla(ws_dest) is None:
                    continue
                datos = buscar_tabla_origen(wb_origen, nombre_hoja, wb_formulas)
                if datos:
                    _volcar_itemizado(ws_dest, datos, calcular_totales, valores_ia_tablas,
                                      columnas_ia_hoja(columnas_ia, nombre_hoja),
                                      (columnas_usuario or {}).get(nombre_hoja))
                else:
                    avisos.append(aviso_tabla_sin_origen(nombre_hoja))
            except Exception:
                avisos.append(aviso_tabla_fallo(nombre_hoja))
        if wb_origen:
            wb_origen.close()
        if wb_formulas:
            wb_formulas.close()

    # openpyxl guarda las fórmulas sin su resultado: que Excel las calcule al
    # abrir, o se verían vacías hasta editar alguna celda.
    wb.calculation.fullCalcOnLoad = True

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()