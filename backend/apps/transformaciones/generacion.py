"""
Servicio de generación del documento final.

Toma el archivo Excel de la plantilla (con su formato: colores, celdas amarillas,
título) y escribe los datos mapeados DENTRO de él, sin romper el diseño. Es el
paso que convierte el mapeo aprobado en un archivo entregable.

Preserva el formato porque abre el archivo original de la plantilla con openpyxl
y solo escribe en las celdas de destino; no reconstruye nada. Las fórmulas,
estilos y demás hojas quedan intactos.

Tres formas de ubicar dónde escribir cada valor, en orden de prioridad:
  1. celda_destino explícita (ej. "B2") → escribe ahí directo.
  2. etiqueta_busqueda (ej. "RUT:") → busca ese texto y escribe en la celda de
     al lado (a la derecha).
  3. Si no hay ninguna, busca el nombre del campo como etiqueta.

Esto hace que funcione con formularios verticales (etiqueta | casilla) sin tener
que configurar celda por celda a mano.
"""
import io

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string, coordinate_to_tuple


def _buscar_etiqueta(ws, texto):
    """
    Busca una celda cuyo texto coincida (aprox.) con `texto` y devuelve la
    coordenada de la celda de al lado (a la derecha), que es donde va el valor
    en un formulario vertical. Devuelve None si no la encuentra.
    """
    objetivo = texto.strip().lower().rstrip(":")
    for fila in ws.iter_rows():
        for celda in fila:
            if celda.value is None:
                continue
            valor = str(celda.value).strip().lower().rstrip(":")
            if valor == objetivo or objetivo in valor:
                # La casilla a rellenar es la de la derecha.
                return celda.row, celda.column + 1
    return None


def generar_documento(ruta_plantilla, mapeos_con_valor):
    """
    ruta_plantilla: ruta al .xlsx de la plantilla (formato destino).
    mapeos_con_valor: lista de dicts con:
        - campo: el CampoPlantilla (para saber dónde escribir)
        - valor: el dato a escribir (ya tomado del origen)

    Devuelve los bytes del Excel generado, listo para guardar/descargar.
    """
    wb = load_workbook(ruta_plantilla)

    for item in mapeos_con_valor:
        campo = item["campo"]
        valor = item["valor"]
        if valor is None:
            continue

        # Elegir la hoja: la del campo si se indicó, si no la activa.
        if campo.hoja_destino and campo.hoja_destino in wb.sheetnames:
            ws = wb[campo.hoja_destino]
        else:
            ws = wb.active

        destino = None

        # 1. Celda exacta.
        if campo.celda_destino:
            destino = coordinate_to_tuple(campo.celda_destino)  # (fila, col)

        # 2. Buscar por etiqueta explícita.
        if destino is None and campo.etiqueta_busqueda:
            destino = _buscar_etiqueta(ws, campo.etiqueta_busqueda)

        # 3. Buscar por el nombre del campo (reemplazando _ por espacio).
        if destino is None:
            destino = _buscar_etiqueta(ws, campo.nombre.replace("_", " "))

        if destino is None:
            # No se encontró dónde escribir este campo; se omite en silencio.
            # (En una versión futura, esto podría reportarse al usuario.)
            continue

        fila, col = destino
        ws.cell(row=fila, column=col, value=valor)

    # Guardar en memoria y devolver los bytes.
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
