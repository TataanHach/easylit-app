"""
Vistas de plantillas.

Mismo aislamiento por organización que las transformaciones. La galería usa el
serializer ligero; el detalle trae los campos anidados para el editor de mapeo.
"""
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.response import Response

from apps.plantillas.models import Plantilla, UnidadCanonica

from .serializers import (
    PlantillaListaSerializer,
    PlantillaSerializer,
    UnidadCanonicaSerializer,
)


class PlantillaViewSet(viewsets.ModelViewSet):
    """
    GET    /api/plantillas/            galería (ligera)
    POST   /api/plantillas/            crear
    GET    /api/plantillas/{id}/       detalle con campos
    GET    /api/plantillas/{id}/hojas/ hojas del Excel de la plantilla
    """

    def get_queryset(self):
        usuario = self.request.user
        if usuario.es_superadmin:
            qs = Plantilla.objects.all()
        else:
            qs = Plantilla.objects.filter(organizacion=usuario.organizacion)

        # No mostrar archivadas salvo que se pidan explícitamente.
        if self.request.query_params.get("archivadas") != "true":
            qs = qs.filter(archivada=False)

        return qs.prefetch_related("campos")

    def get_serializer_class(self):
        if self.action == "list":
            return PlantillaListaSerializer
        return PlantillaSerializer

    def perform_create(self, serializer):
        serializer.save(
            organizacion=self.request.user.organizacion,
            creada_por=self.request.user,
        )

    @action(detail=True, methods=["get"])
    def hojas(self, request, pk=None):
        """
        Devuelve la lista de hojas del Excel de la plantilla, para que el editor
        de campos pueda asignar cada campo a su hoja. Ej: ["Identificación",
        "Experiencia", "Oferta"].

        Si el archivo no es un Excel legible, devuelve lista vacía (no rompe: el
        editor simplemente no mostrará hojas y el campo irá a la hoja activa).
        """
        plantilla = self.get_object()
        hojas = []
        try:
            from openpyxl import load_workbook
            wb = load_workbook(plantilla.archivo.path, read_only=True)
            hojas = wb.sheetnames
            wb.close()
        except Exception:
            hojas = []
        return Response({"hojas": hojas})


class UnidadCanonicaViewSet(viewsets.ModelViewSet):
    """Catálogo de unidades canónicas de la organización."""
    serializer_class = UnidadCanonicaSerializer

    def get_queryset(self):
        usuario = self.request.user
        if usuario.es_superadmin:
            return UnidadCanonica.objects.all()
        return UnidadCanonica.objects.filter(organizacion=usuario.organizacion)

    def perform_create(self, serializer):
        serializer.save(organizacion=self.request.user.organizacion)


class CampoPlantillaViewSet(viewsets.ModelViewSet):
    """
    CRUD de campos de una plantilla.
      GET/POST   /api/campos/           listar (filtrado por ?plantilla=id) / crear
      PATCH/DELETE /api/campos/{id}/    editar / borrar
    """
    from apps.plantillas.api.serializers import CampoPlantillaSerializer
    serializer_class = CampoPlantillaSerializer

    def get_queryset(self):
        from apps.plantillas.models import CampoPlantilla
        usuario = self.request.user
        if usuario.es_superadmin:
            qs = CampoPlantilla.objects.all()
        else:
            qs = CampoPlantilla.objects.filter(plantilla__organizacion=usuario.organizacion)
        plantilla_id = self.request.query_params.get("plantilla")
        if plantilla_id:
            qs = qs.filter(plantilla_id=plantilla_id)
        return qs.select_related("plantilla")